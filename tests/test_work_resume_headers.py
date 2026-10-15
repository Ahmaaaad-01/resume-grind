"""Guard date/location alignment independently of overall resume bounds."""
import json
import sys
import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
from build_work_resume import row
from inspect_resume_pdf import check_resume


class WorkResumeHeaderTests(unittest.TestCase):
    def test_date_is_independently_right_aligned_without_tabs(self):
        doc = Document()
        table = row(doc, "Product Manager", "August 2023 - Current", italic=True)
        self.assertEqual(table.cell(0, 0).text, "Product Manager")
        self.assertEqual(table.cell(0, 1).text, "August 2023 - Current")
        self.assertEqual(table.cell(0, 1).paragraphs[0].alignment, WD_ALIGN_PARAGRAPH.RIGHT)
        self.assertFalse(table._tbl.xpath(".//w:tab"))
        for margin in table._tbl.tblPr.find(qn("w:tblCellMar")):
            self.assertEqual(margin.get(qn("w:w")), "0")

    def test_location_header_is_borderless_and_reads_left_then_right(self):
        doc = Document()
        table = row(doc, "Capital One", "New York, New York", bold=True)
        self.assertEqual([cell.text for cell in table.rows[0].cells], ["Capital One", "New York, New York"])
        for edge in table._tbl.tblPr.find(qn("w:tblBorders")):
            self.assertEqual(edge.get(qn("w:val")), "nil")

    def test_pdf_header_alignment_is_checked_even_when_bounds_pass(self):
        page = {"width_pt": 612, "height_pt": 792,
                "content_bounds": {"left": 36, "right": 576, "top": 24, "bottom": 739},
                "fonts": [{"font": "TimesNewRomanPSMT", "size": 12}],
                "bullet_line_counts": [1, 1],
                "lines": [{"text": "Capital One New York, New York", "x1": 480}]}
        report = {"page_count": 1, "pages": [page]}
        plan = {"bullets": [], "header_rows": [{"left": "Capital One", "right": "New York, New York", "right_edge_pt": 576}]}
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "plan.json"
            path.write_text(json.dumps(plan), encoding="utf-8")
            result = check_resume(report, path)
            self.assertFalse(result["passed"])
            self.assertTrue(any("right margin" in error for error in result["errors"]))
            page["lines"][0]["x1"] = 576.1
            self.assertTrue(check_resume(report, path)["passed"])


if __name__ == "__main__":
    unittest.main()
