"""Build an editable resume from the current profile and a Work-authored draft.

This performs no discovery, scoring, LLM calls, or automatic content trimming.
Render the DOCX to PDF and inspect it before treating the result as finished.
"""
from __future__ import annotations

import argparse
import json
import os
import struct
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

from build_resume_skills import build_catalog, skill_lookup

ROOT = Path(__file__).resolve().parents[1]
BODY_WIDTH = 522.0  # Letter minus 0.5-inch margins and 0.25-inch bullet indent
BODY_SIZE = 12
LINE_HEIGHT = 14.32
KERN_PAIRS = {}
FONT_UNITS = 2048


def load(path):
    return json.loads(Path(path).read_text(encoding="utf-8"))


def register_font(font_dir):
    global KERN_PAIRS, FONT_UNITS
    candidates = [font_dir] if font_dir else [
        Path(os.environ.get("WINDIR", "C:/Windows")) / "Fonts",
        Path("/usr/share/fonts/truetype/msttcorefonts"),
    ]
    for directory in candidates:
        for name in ("times.ttf", "Times_New_Roman.ttf", "Times_New_Roman_Regular.ttf"):
            font = Path(directory) / name
            if font.exists():
                pdfmetrics.registerFont(TTFont("ResumeTNR", str(font)))
                data = font.read_bytes()
                count = struct.unpack_from(">H", data, 4)[0]
                tables = {data[12 + i * 16:16 + i * 16].decode("ascii"):
                          struct.unpack_from(">II", data, 20 + i * 16) for i in range(count)}
                FONT_UNITS = struct.unpack_from(">H", data, tables["head"][0] + 18)[0]
                KERN_PAIRS = {}
                if "kern" in tables:
                    start = tables["kern"][0]
                    subtables = struct.unpack_from(">H", data, start + 2)[0]
                    cursor = start + 4
                    for _ in range(subtables):
                        _, length, coverage = struct.unpack_from(">HHH", data, cursor)
                        if coverage >> 8 == 0 and coverage & 1:
                            pairs = struct.unpack_from(">H", data, cursor + 6)[0]
                            for i in range(pairs):
                                left, right, value = struct.unpack_from(">HHh", data, cursor + 14 + i * 6)
                                KERN_PAIRS[(left, right)] = value
                        cursor += length
                return
    raise ValueError("Times New Roman is required for accurate line measurement. Provide --font-dir; do not substitute fonts.")


def measure(text):
    face = pdfmetrics.getFont("ResumeTNR").face
    glyphs = [face.charToGlyph.get(ord(char), 0) for char in text]
    kern = sum(KERN_PAIRS.get(pair, 0) for pair in zip(glyphs, glyphs[1:]))
    return pdfmetrics.stringWidth(text, "ResumeTNR", BODY_SIZE) + kern * BODY_SIZE / FONT_UNITS


def wrap(text, width=BODY_WIDTH):
    lines = []
    current = ""
    for word in text.split():
        if measure(word) > width:
            raise ValueError(f"Unbreakable text exceeds the line width: {word}")
        proposed = f"{current} {word}".strip()
        if current and measure(proposed) > width:
            lines.append(current)
            current = word
        else:
            current = proposed
    if current:
        lines.append(current)
    return lines


def groups_for_section(profile, bank, index):
    groups = list(profile["sections"][index]["groups"])
    if index == 0:
        groups.extend(bank.get("optional_groups", []))
    return groups


def all_roles(profile, bank):
    return {role["id"]: role for i in range(len(profile["sections"]))
            for group in groups_for_section(profile, bank, i) for role in group["roles"]}


def baseline_draft(profile, bank):
    return {
        "role_bullets": {
            role["id"]: [{"text": bullet["text"], "source_ids": [bullet["id"]],
                          "max_lines": bullet["max_lines"]} for bullet in role["bullets"]]
            for role in all_roles(profile, bank).values()
        },
        "skills": profile["skills"],
        "gaps": [],
    }


def validate(profile, draft, bank, skills_catalog=None):
    roles = all_roles(profile, bank)
    if set(draft.get("role_bullets", {})) != set(roles):
        raise ValueError("Draft must specify every current role ID. Use [] to explicitly omit an optional role.")
    for role_id, bullets in draft["role_bullets"].items():
        allowed = {b["id"] for b in bank["roles"][role_id]["bullets"]}
        if role_id in ("co-servicing-ai", "co-data-products", "surops-founder") and not bullets:
            raise ValueError(f"Keep the current resume's core role: {role_id}")
        for bullet in bullets:
            sources = bullet.get("source_ids", [])
            if not sources or not set(sources) <= allowed:
                raise ValueError(f"Bullet in {role_id} needs source IDs from that same role.")
            if not bullet.get("text", "").strip():
                raise ValueError(f"Empty bullet in {role_id}")
            if not 1 <= int(bullet.get("max_lines", 2)) <= 3:
                raise ValueError("Each bullet must have a line budget between one and three.")
    catalog = skills_catalog if skills_catalog is not None else build_catalog(profile, bank, {"skills": []})
    lookup = skill_lookup(catalog, profile, bank)
    groups = draft.get("skills", profile["skills"])
    if len(groups) != 2:
        raise ValueError("Keep two skill lines, matching the current layout.")
    labels = set()
    selected = set()
    skill_report = []
    for group in groups:
        label = group.get("label", "")
        if (not isinstance(label, str) or not label.strip() or label != label.strip()
                or len(label) > 40 or any(c in label for c in "\n\r\t:")
                or label.casefold() in labels):
            raise ValueError("Skill group labels must be distinct, concise plain text.")
        labels.add(label.casefold())
        items = group.get("items", [])
        if not items:
            raise ValueError("Each skill line must contain supported skills.")
        for item in items:
            if not isinstance(item, str) or item != item.strip() or item.casefold() not in lookup:
                raise ValueError(f"Skill must be in the source-backed catalog: {item}")
            entry = lookup[item.casefold()]
            key = entry["name"].casefold()
            if key in selected:
                raise ValueError(f"Duplicate skill or alias in the resume: {item}")
            selected.add(key)
            skill_report.append({"label": label, "display": item, "canonical": entry["name"],
                                 "basis": entry["basis"], "source_ids": entry["source_ids"]})
    return skill_report


def font(run, *, bold=False, italic=False, size=12, name="Times New Roman"):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor(0, 0, 0)
    props = run._element.get_or_add_rPr()
    fonts = props.find(qn("w:rFonts"))
    if fonts is not None:
        for key in ("ascii", "hAnsi", "eastAsia", "cs"):
            fonts.set(qn(f"w:{key}"), name)
    if name == "Times New Roman":
        kerning = OxmlElement("w:kern")
        kerning.set(qn("w:val"), "24")
        props.append(kerning)


def paragraph(doc, *, before=0, after=0.5, line=LINE_HEIGHT, keep=False):
    p = doc.add_paragraph()
    fmt = p.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = Pt(line)
    fmt.keep_together = True
    fmt.keep_with_next = keep
    fmt.widow_control = True
    return p


def border(p, side="top", space=6):
    props = p._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    edge = OxmlElement(f"w:{side}")
    for key, value in {"val": "single", "sz": "4", "space": str(space), "color": "808080"}.items():
        edge.set(qn(f"w:{key}"), value)
    borders.append(edge)
    props.append(borders)


def row(doc, left, right="", *, bold=False, italic=False):
    if not right:
        p = paragraph(doc, keep=True)
        font(p.add_run(left), bold=bold, italic=italic)
        return p
    # Some DOCX previews collapse custom tab stops. Simple, inline header rows
    # give each value explicit alignment; all achievement text stays in paragraphs.
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    widths = (360, 180)
    for column, width in zip(table.columns, widths):
        column.width = Pt(width)
    props = table._tbl.tblPr
    width = props.find(qn("w:tblW"))
    width.set(qn("w:w"), "10800")
    width.set(qn("w:type"), "dxa")
    margins = OxmlElement("w:tblCellMar")
    for side in ("top", "left", "bottom", "right"):
        edge = OxmlElement(f"w:{side}")
        edge.set(qn("w:w"), "0")
        edge.set(qn("w:type"), "dxa")
        margins.append(edge)
    props.append(margins)
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        edge = OxmlElement(f"w:{side}")
        edge.set(qn("w:val"), "nil")
        borders.append(edge)
    props.append(borders)
    indent = OxmlElement("w:tblInd")
    indent.set(qn("w:w"), "0")
    indent.set(qn("w:type"), "dxa")
    props.append(indent)
    tr = table.rows[0]
    tr.height = Pt(LINE_HEIGHT + 0.5)
    tr.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
    tr._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
    for cell, width, text, alignment, is_italic in zip(
            tr.cells, widths, (left, right),
            (WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.RIGHT), (italic, False)):
        cell.width = Pt(width)
        p = cell.paragraphs[0]
        fmt = p.paragraph_format
        fmt.left_indent = fmt.right_indent = fmt.first_line_indent = Pt(0)
        fmt.space_before = fmt.space_after = Pt(0)
        fmt.line_spacing = Pt(LINE_HEIGHT)
        fmt.keep_together = fmt.keep_with_next = True
        p.alignment = alignment
        font(p.add_run(text), bold=bold, italic=is_italic)
    return table


def heading(doc, label, *, first=False, ruled=True):
    p = paragraph(doc, before=0 if first else (9 if ruled else 0), keep=True)
    font(p.add_run(label), bold=True)
    if ruled:
        border(p)
    return p


def bullet(doc, text):
    lines = wrap(text)
    p = paragraph(doc)
    p.paragraph_format.left_indent = Pt(18)
    p.paragraph_format.first_line_indent = Pt(-18)
    p.paragraph_format.tab_stops.add_tab_stop(Pt(18))
    font(p.add_run("\u2022\t"), name="Arial")
    for i, line in enumerate(lines):
        if i:
            p.add_run().add_break()
        font(p.add_run(line))
    return lines


def build(profile, draft, output, bank, skills_catalog=None):
    skill_report = validate(profile, draft, bank, skills_catalog)
    report = {"body_font": "Times New Roman", "body_size": BODY_SIZE,
              "bullet_width_pt": BODY_WIDTH, "bullets": [], "warnings": [],
              "requires_pdf_render": True, "factual_review_required": True,
              "skills": skill_report, "header_rows": []}
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Inches(8.5)
    sec.page_height = Inches(11)
    sec.top_margin = Pt(21.4)
    sec.bottom_margin = Pt(36)
    sec.left_margin = sec.right_margin = Pt(36)
    sec.header_distance = sec.footer_distance = Pt(0)
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(0.5)
    normal.paragraph_format.line_spacing = Pt(LINE_HEIGHT)
    title_props = doc.styles["Title"]._element.pPr
    if title_props is not None:
        for item in list(title_props):
            if item.tag == qn("w:pBdr"):
                title_props.remove(item)
    doc.core_properties.title = "Resume"
    doc.core_properties.author = profile["contact"]["name"]
    doc.core_properties.subject = ""
    p = paragraph(doc, after=0, line=23.62, keep=True)
    p.style = doc.styles["Title"]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    font(p.add_run(profile["contact"]["name"]), size=20)
    p = paragraph(doc, after=14.19, keep=True)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    font(p.add_run(" | ".join(profile["contact"][key] for key in ("location", "phone", "email"))))
    for section_index, section in enumerate(profile["sections"]):
        heading(doc, section["heading"], first=section_index == 0)
        for group in groups_for_section(profile, bank, section_index):
            visible = [r for r in group["roles"] if draft["role_bullets"][r["id"]]]
            if not visible:
                continue
            row(doc, group["company"], group["location"], bold=True)
            report["header_rows"].append({"left": group["company"], "right": group["location"], "right_edge_pt": 576})
            for role in visible:
                row(doc, role["title"], role["date_display"], italic=True)
                if role["date_display"]:
                    report["header_rows"].append({"left": role["title"], "right": role["date_display"], "right_edge_pt": 576})
                for entry in draft["role_bullets"][role["id"]]:
                    lines = wrap(entry["text"])
                    budget = int(entry.get("max_lines", 2))
                    if len(lines) > budget:
                        raise ValueError(f"{role['id']}: bullet has {len(lines)} lines, budget {budget}. Shorten the text: {entry['text']}")
                    widths = [round(measure(line), 2) for line in lines]
                    if len(lines) > 1 and widths[-1] < BODY_WIDTH * 0.30:
                        report["warnings"].append(f"{role['id']}: short final line; rewrite to avoid a dangling fragment")
                    bullet(doc, entry["text"])
                    report["bullets"].append({"role_id": role["id"], "source_ids": entry["source_ids"],
                                              "line_count": len(lines), "max_lines": budget,
                                              "line_widths_pt": widths})
    heading(doc, "SKILLS")
    skill_groups = draft.get("skills", profile["skills"])
    if len(skill_groups) != 2 or len({g['label'] for g in skill_groups}) != 2:
        raise ValueError("Keep two distinct skill lines, matching the current layout.")
    for group in skill_groups:
        text = f"{group['label']}: " + ", ".join(group["items"])
        if len(wrap(text)) > 1:
            raise ValueError(f"{group['label']} skills exceed one line; select fewer relevant skills.")
        lines = bullet(doc, text)
        report.setdefault("skill_lines", []).append({"label": group["label"],
                                                     "line_count": len(lines), "width_pt": round(measure(text), 2)})
    heading(doc, "EDUCATION", ruled=False)
    edu = profile["education"]
    row(doc, edu["school"], edu["location"], bold=True)
    report["header_rows"].append({"left": edu["school"], "right": edu["location"], "right_edge_pt": 576})
    p = paragraph(doc)
    font(p.add_run(edu["degree"]), italic=True)
    p = paragraph(doc)
    font(p.add_run("Certifications: "), bold=True)
    font(p.add_run(edu["certification"]))
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    report["total_experience_bullet_lines"] = sum(b["line_count"] for b in report["bullets"])
    output.with_suffix(".layout.json").write_text(json.dumps(report, indent=2), encoding="utf-8")
    print(f"Saved {output}; bullet lines: {report['total_experience_bullet_lines']}. Render and inspect before delivery.")
    for warning in report["warnings"]:
        print("Warning: " + warning)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--profile", type=Path, default=ROOT / "data/resume_workflow/current_profile.json")
    parser.add_argument("--draft", type=Path)
    parser.add_argument("--bank", type=Path, default=ROOT / "data/resume_workflow/approved_bank.json")
    parser.add_argument("--skills-catalog", type=Path, default=ROOT / "data/resume_workflow/skills_catalog.json")
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--font-dir", type=Path)
    parser.add_argument("--export-draft", type=Path)
    args = parser.parse_args()
    profile = load(args.profile)
    bank = load(args.bank)
    draft = load(args.draft) if args.draft else baseline_draft(profile, bank)
    register_font(args.font_dir)
    catalog = load(args.skills_catalog) if args.skills_catalog.exists() else None
    build(profile, draft, args.output, bank, catalog)
    if args.export_draft:
        args.export_draft.parent.mkdir(parents=True, exist_ok=True)
        args.export_draft.write_text(json.dumps(draft, indent=2), encoding="utf-8")


if __name__ == "__main__":
    main()
